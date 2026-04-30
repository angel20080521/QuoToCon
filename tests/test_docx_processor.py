"""
tests/test_docx_processor.py
Unit tests for utils/docx_processor.py
"""

import io
import os
import sys
import pytest
from datetime import datetime

# Make the project root importable regardless of where pytest is invoked from.
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from docx import Document
from utils.docx_processor import (
    extract_data_from_quotation_xlsx,
    fill_template,
    generate_output_filename,
    _amount_to_chinese,
    _format_amount_str,
    _extract_kv_from_row,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_xlsx(rows: list[list], sheet_name: str = 'Sheet1') -> str:
    """Write a temporary .xlsx with given rows and return its path."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    tmp = os.path.join(os.path.dirname(__file__), '_tmp_quotation.xlsx')
    wb.save(tmp)
    return tmp


def _make_template_docx(paragraphs: list[str]) -> str:
    """Write a temporary template .docx and return its path."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    tmp = os.path.join(os.path.dirname(__file__), '_tmp_template.docx')
    doc.save(tmp)
    return tmp


def _cleanup(*paths):
    for p in paths:
        try:
            os.remove(p)
        except OSError:
            pass


# ---------------------------------------------------------------------------
# extract_data_from_quotation_xlsx
# ---------------------------------------------------------------------------

class TestExtractDataXlsx:

    def test_two_column_table_fullwidth_key(self, tmp_path):
        """Standard two-column table with full-width field names."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['客户名称：', '某某有限公司'])
        ws.append(['报价编号：', 'QT-2024-001'])
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        assert data['客户名称'] == '某某有限公司'
        assert data['报价编号'] == 'QT-2024-001'

    def test_two_column_table_key_without_colon(self, tmp_path):
        """Full-width field name without trailing colon."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['项目名称', '智慧园区项目'])
        ws.append(['签约日期', '2024-05-01'])
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        assert data['项目名称'] == '智慧园区项目'
        assert data['签约日期'] == '2024-05-01'

    def test_single_cell_fullwidth_colon_pattern(self, tmp_path):
        """Single cell with full-width 'key：value' format."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['合计金额：46,000 元'])
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        assert data.get('合计金额') == '46,000 元'

    def test_empty_rows_ignored(self, tmp_path):
        """Completely empty rows are skipped."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([None, None])
        ws.append(['客户名称：', '测试公司'])
        ws.append([None])
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        assert data['客户名称'] == '测试公司'
        assert len([k for k in data if not k.startswith('_')]) == 1

    def test_multi_column_rows_stored_as_private(self, tmp_path):
        """Rows with 3+ non-empty cells become private table metadata."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['型号', '数量', '单价'])
        ws.append(['TypeA', '10', '500'])
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        private_keys = [k for k in data if k.startswith('_')]
        assert len(private_keys) > 0

    def test_first_sheet_value_wins(self, tmp_path):
        """When the same key appears on multiple sheets, first wins."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = 'Sheet1'
        ws1.append(['客户名称：', '第一页公司'])
        ws2 = wb.create_sheet('Sheet2')
        ws2.append(['客户名称：', '第二页公司'])
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        assert data['客户名称'] == '第一页公司'

    def test_numeric_value_converted_to_string(self, tmp_path):
        """Numeric cell values are stringified."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['合计金额：', 46000])
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        assert data['合计金额'] == '46000'

    def test_single_cell_halfwidth_colon_not_matched(self, tmp_path):
        """Half-width colon in a single cell is NOT extracted (full-width only)."""
        q = str(tmp_path / 'q.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['客户名称: 某公司'])   # half-width colon
        wb.save(q)

        data = extract_data_from_quotation_xlsx(q)
        assert '客户名称' not in data


# ---------------------------------------------------------------------------
# fill_template  (unchanged – kept for regression)
# ---------------------------------------------------------------------------

class TestFillTemplate:

    def test_basic_replacement(self, tmp_path):
        tmpl = str(tmp_path / 't.docx')
        out = str(tmp_path / 'out.docx')
        doc = Document()
        doc.add_paragraph('甲方：{{客户名称}}')
        doc.add_paragraph('编号：{{报价编号}}')
        doc.save(tmpl)

        data = {'客户名称': '某某公司', '报价编号': 'QT-001'}
        fill_template(tmpl, data, out)

        result = Document(out)
        texts = [p.text for p in result.paragraphs]
        assert '甲方：某某公司' in texts
        assert '编号：QT-001' in texts

    def test_unmatched_placeholder_retained(self, tmp_path):
        tmpl = str(tmp_path / 't.docx')
        out = str(tmp_path / 'out.docx')
        doc = Document()
        doc.add_paragraph('合同金额：{{合计金额}}')
        doc.save(tmpl)

        fill_template(tmpl, {}, out)

        result = Document(out)
        texts = [p.text for p in result.paragraphs]
        assert any('{{合计金额}}' in t for t in texts)

    def test_private_keys_not_replaced(self, tmp_path):
        tmpl = str(tmp_path / 't.docx')
        out = str(tmp_path / 'out.docx')
        doc = Document()
        doc.add_paragraph('{{_internal}}')
        doc.save(tmpl)

        fill_template(tmpl, {'_internal': '敏感数据'}, out)

        result = Document(out)
        texts = [p.text for p in result.paragraphs]
        assert any('{{_internal}}' in t for t in texts)

    def test_replacement_in_table_cell(self, tmp_path):
        tmpl = str(tmp_path / 't.docx')
        out = str(tmp_path / 'out.docx')
        doc = Document()
        tbl = doc.add_table(rows=1, cols=2)
        tbl.rows[0].cells[0].text = '金额'
        tbl.rows[0].cells[1].text = '{{合计金额}}'
        doc.save(tmpl)

        fill_template(tmpl, {'合计金额': '99,000 元'}, out)

        result = Document(out)
        cell_text = result.tables[0].rows[0].cells[1].text
        assert cell_text == '99,000 元'

    def test_output_directory_created_if_missing(self, tmp_path):
        tmpl = str(tmp_path / 't.docx')
        doc = Document()
        doc.add_paragraph('test')
        doc.save(tmpl)

        nested_out = str(tmp_path / 'subdir' / 'deep' / 'out.docx')
        fill_template(tmpl, {}, nested_out)
        assert os.path.isfile(nested_out)


# ---------------------------------------------------------------------------
# generate_output_filename
# ---------------------------------------------------------------------------

class TestGenerateOutputFilename:

    def test_uses_quotation_number_first(self):
        data = {'报价编号': 'QT-2024-001', '客户名称': '某公司'}
        name = generate_output_filename(data)
        assert name.startswith('QT-2024-001_')
        assert name.endswith('.docx')

    def test_falls_back_to_project_name(self):
        data = {'项目名称': '智慧园区'}
        name = generate_output_filename(data)
        assert name.startswith('智慧园区_')

    def test_falls_back_to_default(self):
        name = generate_output_filename({})
        assert name.startswith('合同_')

    def test_invalid_chars_sanitised(self):
        data = {'报价编号': 'QT/2024\\001:test'}
        name = generate_output_filename(data)
        assert '/' not in name
        assert '\\' not in name
        assert ':' not in name

    def test_timestamp_suffix_format(self):
        data = {'报价编号': 'QT-001'}
        name = generate_output_filename(data)
        stem = name[:-5]  # strip .docx
        parts = stem.rsplit('_', 2)
        assert len(parts) == 3
        date_part, time_part = parts[1], parts[2]
        datetime.strptime(date_part + time_part, '%Y%m%d%H%M%S')

    def test_name_truncated_to_50_chars(self):
        long_name = '甲' * 100
        data = {'报价编号': long_name}
        name = generate_output_filename(data)
        stem = name[:-5]
        name_part = stem.rsplit('_', 2)[0]
        assert len(name_part) <= 50

    def test_private_keys_not_used_for_name(self):
        data = {'_internal': 'secret', '客户名称': '某公司'}
        name = generate_output_filename(data)
        assert name.startswith('某公司_')


# ---------------------------------------------------------------------------
# _amount_to_chinese
# ---------------------------------------------------------------------------

class TestAmountToChinese:

    def test_zero(self):
        assert _amount_to_chinese(0) == '零元整'

    def test_whole_yuan(self):
        assert _amount_to_chinese(1) == '壹元整'
        assert _amount_to_chinese(10) == '壹拾元整'
        assert _amount_to_chinese(100) == '壹佰元整'
        assert _amount_to_chinese(1000) == '壹仟元整'
        assert _amount_to_chinese(10000) == '壹万元整'
        assert _amount_to_chinese(100000) == '壹拾万元整'

    def test_with_fen_and_jiao(self):
        assert _amount_to_chinese(0.01) == '零元零壹分'
        assert _amount_to_chinese(0.10) == '零元壹角整'
        assert _amount_to_chinese(1.23) == '壹元贰角叁分'

    def test_typical_contract_amount(self):
        result = _amount_to_chinese(12345.67)
        assert '壹万' in result
        assert '贰仟' in result
        assert '叁佰' in result
        assert '肆拾' in result
        assert '伍元' in result
        assert '陆角' in result
        assert '柒分' in result

    def test_zero_middle_group(self):
        # 10001 → 壹万零壹元整
        result = _amount_to_chinese(10001)
        assert '零' in result
        assert '壹元' in result

    def test_negative(self):
        assert _amount_to_chinese(-100) == '负壹佰元整'

    def test_none_returns_empty(self):
        assert _amount_to_chinese(None) == ''

    def test_invalid_returns_str(self):
        assert _amount_to_chinese('abc') == 'abc'


# ---------------------------------------------------------------------------
# _format_amount_str
# ---------------------------------------------------------------------------

class TestFormatAmountStr:

    def test_integer(self):
        assert _format_amount_str(1000) == '1000.00'

    def test_float(self):
        assert _format_amount_str(12345.678) == '12345.68'  # rounded

    def test_none(self):
        assert _format_amount_str(None) == ''

    def test_string_number(self):
        assert _format_amount_str('500') == '500.00'


# ---------------------------------------------------------------------------
# _extract_kv_from_row
# ---------------------------------------------------------------------------

class TestExtractKvFromRow:

    def test_single_label_value(self):
        result = _extract_kv_from_row(['客户名称：', '某公司'])
        assert result == {'客户名称': '某公司'}

    def test_multiple_labels_per_row(self):
        result = _extract_kv_from_row(
            ['客户名称：', '某公司', '', '', '联系人：', '张三', '收款方式：', '月结30天']
        )
        assert result['客户名称'] == '某公司'
        assert result['联系人'] == '张三'
        assert result['收款方式'] == '月结30天'

    def test_label_without_value(self):
        result = _extract_kv_from_row(['报价单号码：'])
        assert result.get('报价单号码') == ''

    def test_adjacent_labels(self):
        # Two labels in a row: first gets empty value
        result = _extract_kv_from_row(['字段A：', '字段B：', '值B'])
        assert result['字段A'] == ''
        assert result['字段B'] == '值B'

    def test_no_labels_returns_empty(self):
        assert _extract_kv_from_row(['型号', '数量', '单价']) == {}

    def test_halfwidth_colon_not_matched(self):
        result = _extract_kv_from_row(['客户名称:', '某公司'])
        assert '客户名称' not in result


# ---------------------------------------------------------------------------
# extract_data_from_quotation_xlsx – new capabilities
# ---------------------------------------------------------------------------

class TestExtractDataXlsxNewCapabilities:

    def _make_xlsx(self, rows, tmp_path, name='q.xlsx'):
        path = str(tmp_path / name)
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        wb.save(path)
        return path

    def test_multi_field_row_extraction(self, tmp_path):
        """Multi-field rows: several 'label：' + value pairs on one row."""
        q = self._make_xlsx([
            ['客户名称：', '测试公司', None, None, '联系人：', '李四', '发票类型：', '增值税-专用发票'],
        ], tmp_path)
        data = extract_data_from_quotation_xlsx(q)
        assert data['客户名称'] == '测试公司'
        assert data['联系人'] == '李四'
        assert data['发票类型'] == '增值税-专用发票'

    def test_fblx_mapping_zhuanyong(self, tmp_path):
        """发票类型 '增值税-专用发票' maps to fblx='2'."""
        q = self._make_xlsx([['发票类型：', '增值税-专用发票']], tmp_path)
        data = extract_data_from_quotation_xlsx(q)
        assert data.get('fblx') == '2'

    def test_fblx_mapping_putong(self, tmp_path):
        """发票类型 '普通发票' maps to fblx='1'."""
        q = self._make_xlsx([['发票类型：', '普通发票']], tmp_path)
        data = extract_data_from_quotation_xlsx(q)
        assert data.get('fblx') == '1'

    def test_product_table_extraction(self, tmp_path):
        """Product table is detected and amounts calculated correctly."""
        rows = [
            ['客户名称：', '客户A'],
            # product table
            ['编号', '商品种类', '商品名称', '产品描述', '数量', '单位',
             '单价（含税）', '单价（未税）', '税率', '金额(未税)', '税额', '金额(含税）'],
            [1, '软硬件商品', '路由器', 'AC-1000', 2, '台', 1130, None, '13%', None, None, None],
            [2, '安装调试服务', '部署服务', '安装配置', 1, '次', 106, None, '6%', None, None, None],
        ]
        q = self._make_xlsx(rows, tmp_path)
        data = extract_data_from_quotation_xlsx(q)

        # Hardware b1* fields
        assert data.get('b1bh') == '1'
        assert data.get('b1spxh') == '路由器'
        assert data.get('b1spmx') == 'AC-1000'
        assert data.get('b1sl') == '2'
        assert data.get('b1dw') == '台'
        assert data.get('b1dj') == '1130.00'

        # Hardware amounts (1130 * 2 = 2260 incl. tax)
        hw_hs = float(data.get('b1zj', '0'))
        assert abs(hw_hs - 2260.0) < 0.01

        # Service b2* fields
        assert data.get('b2bh') == '2'
        assert data.get('b2spxh') == '部署服务'

        # Amount placeholders generated
        assert '合同总额大写' in data
        assert '商品总额大写' in data
        assert '服务总额大写' in data

    def test_amount_to_chinese_in_output(self, tmp_path):
        """Chinese numeral conversion is applied to extracted amounts."""
        rows = [
            ['编号', '商品种类', '商品名称', '产品描述', '数量', '单位',
             '单价（含税）', '单价（未税）', '税率', '金额(未税)', '税额', '金额(含税）'],
            [1, '软硬件商品', '交换机', '48口', 1, '台', 1130, None, '13%', None, None, None],
        ]
        q = self._make_xlsx(rows, tmp_path)
        data = extract_data_from_quotation_xlsx(q)
        # 1130 incl. tax → b1zj should be '1130.00'
        assert data.get('b1zj') == '1130.00'
        # 大写 should contain Chinese numerals
        daxie = data.get('合同总额大写', '')
        assert '壹' in daxie or '仟' in daxie

    def test_product_table_does_not_pollute_table_rows(self, tmp_path):
        """Product-table rows are NOT stored under _sheet_N_table_headers."""
        rows = [
            ['编号', '商品种类', '商品名称', '产品描述', '数量', '单位',
             '单价（含税）', '单价（未税）', '税率', '金额(未税)', '税额', '金额(含税）'],
            [1, '软硬件商品', '设备', '规格A', 1, '台', 1000, None, '13%', None, None, None],
        ]
        q = self._make_xlsx(rows, tmp_path)
        data = extract_data_from_quotation_xlsx(q)
        assert '_sheet_0_table_headers' not in data


# ---------------------------------------------------------------------------
# fill_template – multi-product row expansion
# ---------------------------------------------------------------------------

class TestFillTemplateMultiProduct:

    def test_single_hardware_row_replaced(self, tmp_path):
        """Single hardware product: b1* placeholders are replaced normally."""
        tmpl = str(tmp_path / 't.docx')
        out = str(tmp_path / 'out.docx')
        doc = Document()
        tbl = doc.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text = '{{b1bh}}'
        tbl.rows[0].cells[1].text = '{{b1spxh}}'
        tbl.rows[1].cells[0].text = '总计'
        tbl.rows[1].cells[1].text = '{{b1zj}}'
        doc.save(tmpl)

        data = {
            'b1bh': '1', 'b1spxh': '路由器', 'b1spmx': '规格A',
            'b1sl': '2', 'b1dw': '台', 'b1dj': '1130.00',
            'b1wsje': '2000.00', 'b1hsje': '2260.00', 'b1zj': '2260.00',
            '_sheet_0_hardware_products': [
                {'bh': '1', 'spmc': '路由器', 'mssm': '规格A',
                 'sl': '2', 'dw': '台', 'dj_hs': 1130.0,
                 'je_ws': 2000.0, 'je_hs': 2260.0},
            ],
            '_sheet_0_service_products': [],
        }
        fill_template(tmpl, data, out)

        result = Document(out)
        cell_texts = [cell.text for row in result.tables[0].rows for cell in row.cells]
        assert '1' in cell_texts
        assert '路由器' in cell_texts
        assert '2260.00' in cell_texts

    def test_multiple_hardware_rows_expanded(self, tmp_path):
        """Two hardware products: an extra row is inserted for the second one."""
        tmpl = str(tmp_path / 't.docx')
        out = str(tmp_path / 'out.docx')
        doc = Document()
        tbl = doc.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text = '{{b1bh}}'
        tbl.rows[0].cells[1].text = '{{b1spxh}}'
        tbl.rows[1].cells[0].text = '总计'
        tbl.rows[1].cells[1].text = '{{b1zj}}'
        doc.save(tmpl)

        data = {
            'b1bh': '1', 'b1spxh': '路由器', 'b1spmx': '',
            'b1sl': '1', 'b1dw': '台', 'b1dj': '1000.00',
            'b1wsje': '', 'b1hsje': '', 'b1zj': '2260.00',
            '_sheet_0_hardware_products': [
                {'bh': '1', 'spmc': '路由器', 'mssm': '', 'sl': '1', 'dw': '台',
                 'dj_hs': 1000, 'je_ws': None, 'je_hs': None},
                {'bh': '2', 'spmc': '交换机', 'mssm': '', 'sl': '2', 'dw': '台',
                 'dj_hs': 630, 'je_ws': None, 'je_hs': None},
            ],
            '_sheet_0_service_products': [],
        }
        fill_template(tmpl, data, out)

        result = Document(out)
        all_texts = [cell.text for row in result.tables[0].rows for cell in row.cells]
        assert '路由器' in all_texts
        assert '交换机' in all_texts
        # Three rows: product 1, product 2, totals
        assert len(result.tables[0].rows) == 3
