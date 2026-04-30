"""
tests/test_app.py
Integration tests for the Flask routes in app.py
"""

import io
import os
import sys
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from docx import Document
import app as flask_app


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture()
def client(tmp_path):
    flask_app.app.config['TESTING'] = True
    flask_app.app.config['UPLOAD_FOLDER'] = str(tmp_path / 'uploads')
    flask_app.app.config['OUTPUT_FOLDER'] = str(tmp_path / 'outputs')
    os.makedirs(flask_app.app.config['UPLOAD_FOLDER'], exist_ok=True)
    os.makedirs(flask_app.app.config['OUTPUT_FOLDER'], exist_ok=True)
    with flask_app.app.test_client() as c:
        yield c


def _make_xlsx_bytes(rows: list[list]) -> bytes:
    """Return the raw bytes of a minimal .xlsx with the given rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    """Return the raw bytes of a minimal .docx with the given paragraphs."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# GET /
# ---------------------------------------------------------------------------

class TestIndex:

    def test_200_and_form_present(self, client):
        rv = client.get('/')
        assert rv.status_code == 200
        assert b'<form' in rv.data

    def test_html_charset(self, client):
        rv = client.get('/')
        assert b'charset' in rv.data.lower()

    def test_xlsx_accept_shown_for_quotation(self, client):
        rv = client.get('/')
        assert b'.xlsx' in rv.data


# ---------------------------------------------------------------------------
# POST /process
# ---------------------------------------------------------------------------

class TestProcess:

    def test_missing_both_files_redirects(self, client):
        rv = client.post('/process', data={})
        assert rv.status_code == 302

    def test_missing_template_redirects(self, client):
        rv = client.post('/process', data={
            'quotation': (io.BytesIO(_make_xlsx_bytes([['客户名称：', '测试公司']])), 'q.xlsx'),
        }, content_type='multipart/form-data')
        assert rv.status_code == 302

    def test_wrong_quotation_extension_redirects(self, client):
        """Uploading a .docx as quotation should be rejected."""
        rv = client.post('/process', data={
            'quotation': (io.BytesIO(_make_docx_bytes(['客户名称：测试公司'])), 'q.docx'),
            'template': (io.BytesIO(_make_docx_bytes(['{{客户名称}}'])), 't.docx'),
        }, content_type='multipart/form-data')
        assert rv.status_code == 302

    def test_wrong_template_extension_redirects(self, client):
        """Uploading a non-.docx file as template should be rejected."""
        rv = client.post('/process', data={
            'quotation': (io.BytesIO(_make_xlsx_bytes([['客户名称：', '测试公司']])), 'q.xlsx'),
            'template': (io.BytesIO(b'not a docx'), 't.txt'),
        }, content_type='multipart/form-data')
        assert rv.status_code == 302

    def test_successful_process_returns_200_with_download_link(self, client):
        quotation_bytes = _make_xlsx_bytes([
            ['客户名称：', '某某公司'],
            ['报价编号：', 'QT-001'],
        ])
        template_bytes = _make_docx_bytes(['甲方：{{客户名称}}', '编号：{{报价编号}}'])

        rv = client.post('/process', data={
            'quotation': (io.BytesIO(quotation_bytes), 'quotation.xlsx'),
            'template': (io.BytesIO(template_bytes), 'template.docx'),
        }, content_type='multipart/form-data')

        assert rv.status_code == 200
        assert b'download' in rv.data.lower()
        assert '某某公司'.encode() in rv.data or b'QT-001' in rv.data

    def test_extracted_fields_shown_in_result(self, client):
        quotation_bytes = _make_xlsx_bytes([['客户名称：', '示例客户']])
        template_bytes = _make_docx_bytes(['合同甲方：{{客户名称}}'])

        rv = client.post('/process', data={
            'quotation': (io.BytesIO(quotation_bytes), 'q.xlsx'),
            'template': (io.BytesIO(template_bytes), 't.docx'),
        }, content_type='multipart/form-data')

        assert rv.status_code == 200
        assert '客户名称'.encode() in rv.data
        assert '示例客户'.encode() in rv.data


# ---------------------------------------------------------------------------
# GET /download/<filename>
# ---------------------------------------------------------------------------

class TestDownload:

    def test_nonexistent_file_returns_404(self, client):
        rv = client.get('/download/no_such_file.docx')
        assert rv.status_code == 404

    def test_path_traversal_blocked(self, client):
        rv = client.get('/download/../app.py')
        assert rv.status_code in (400, 404)

    def test_download_existing_file(self, client, tmp_path):
        output_dir = flask_app.app.config['OUTPUT_FOLDER']
        fname = '合同_20240501_120000.docx'
        fpath = os.path.join(output_dir, fname)
        doc = Document()
        doc.add_paragraph('test')
        doc.save(fpath)

        rv = client.get(f'/download/{fname}')
        assert rv.status_code == 200
        assert rv.headers['Content-Disposition']
